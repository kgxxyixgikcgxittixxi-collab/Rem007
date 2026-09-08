#!/usr/bin/env python3
# -*- coding: utf-8 -*-
#
# MẪU SINH BÀI THUYẾT TRÌNH CHUẨN — BẢO VỆ MÔI TRƯỜNG
#
# Dùng:  python3 bai_thuyet_trinh_moi_truong.py [path/output.pptx]
#
# Kế hoạch: python-pptx (xem requirements.txt; ARM cài lxml qua apt trước)
#
# Mạng lưới mẫu này là "khuôn chuẩn": 16:9, bảng màu
# thiên nhiên, tiêu đề/header/bullet đồng bộ, đúng cấu trúc OOXML.
"""Tạo bài thuyết trình 'BẢO VỆ MÔI TRƯỜNG' — 16:9, đúng chuẩn OOXML."""
import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as _RGBColor


def RGBColor(*a):
    if len(a) == 1:
        v = a[0]
        return _RGBColor((v >> 16) & 0xFF, (v >> 8) & 0xFF, v & 0xFF)
    return _RGBColor(*a)
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.expanduser("~"), "Downloads", "presentation_moi_truong.pptx")

# ── palette thiên nhiên ──
BG      = RGBColor(0xEAF4EA)   # xanh lá rất nhạt
GREEN_D = RGBColor(0x14532B)   # xanh đậm
GREEN_M = RGBColor(0x1E6B37)
GREEN_L = RGBColor(0x37994A)
GOLD    = RGBColor(0xF2B705)
WHITE   = RGBColor(0xFFFFFF)
TEXT    = RGBColor(0x14301A)
GRAY    = RGBColor(0x5B6B5E)

FONT = "Segoe UI"  # có sẵn trên Windows/Office, đổi fallback nhẹ

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def slide():
    return prs.slides.add_slide(BLANK)


def bg(s, color=BG):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def shape(s, kind, x, y, w, h, color, line=None, alpha=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if color is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = color
        if alpha is not None:
            sh.fill.transparency = alpha
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def text(s, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    """lines: list of (str, size, bold, color, [space_before]) """
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for it in lines:
        txt, sz, bd, col = it[0], it[1], it[2], it[3]
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_before = Pt(it[4] if len(it) > 4 else 0)
        p.space_after = Pt(0)
        run = p.add_run(); run.text = txt
        f = run.font
        f.name = FONT; f.size = Pt(sz); f.bold = bd; f.color.rgb = col
    return tb


def header(s, label, title):
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, 1.05, GREEN_D)
    shape(s, MSO_SHAPE.RECTANGLE, 0, 1.05, W, 0.06, GOLD)
    text(s, 0.6, 0.18, 11.5, 0.5, [(label.upper(), 12, True, GREEN_L)])
    text(s, 0.6, 0.52, 11.5, 0.55, [(title, 30, True, WHITE)])


def chip(s, x, y, w, h, num, label, sub):
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, alpha=0.0, line=GREEN_L)
    # vì background trắng kém tương phản => nền xanh nhạt
    shp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, GREEN_M)
    c = shape(s, MSO_SHAPE.OVAL, x + 0.25, y + 0.25, 0.8, 0.8, GOLD)
    text(s, x + 0.3, y + 0.37, 0.8, 0.5, [(num, 20, True, GREEN_D)], PP_ALIGN.CENTER)
    text(s, x + 1.25, y + 0.28, w - 1.5, h - 0.5, [
        (label, 18, True, WHITE),
        (sub, 12, False, RGBColor(0xE8F5E9))])
    return shp


def bullets(s, x, y, w, h, items):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for i, (b, d) in enumerate(items):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_before = Pt(10); p.space_after = Pt(0)
        r1 = p.add_run(); r1.text = "●  "
        r1.font.name = FONT; r1.font.size = Pt(16); r1.font.bold = True; r1.font.color.rgb = GREEN_L
        r2 = p.add_run(); r2.text = b
        r2.font.name = FONT; r2.font.size = Pt(16); r2.font.bold = True; r2.font.color.rgb = TEXT
        if d:
            p2 = tf.add_paragraph(); p2.space_before = Pt(2)
            r3 = p2.add_run(); r3.text = "      " + d
            r3.font.name = FONT; r3.font.size = Pt(13); r3.font.color.rgb = GRAY
    return tb


# ══════════ 1. SLIDE BÌA ══════════
s = slide()
bg(s, GREEN_D)
for cx, cy, d, al in [(11.2, -0.8, 3.4, 0.55), (11.9, 4.6, 2.6, 0.45), (-1.2, 5.4, 3.0, 0.4)]:
    shape(s, MSO_SHAPE.OVAL, cx, cy, d, d, GREEN_M, alpha=al)
shape(s, MSO_SHAPE.OVAL, 0.6, 6.3, 1.5, 1.5, GOLD, alpha=0.35)
shape(s, MSO_SHAPE.RECTANGLE, 0, 5.9, 4.2, 0.09, GOLD)
text(s, 1.0, 2.15, 11.3, 1.1, [("BẢO VỆ MÔI TRƯỜNG", 60, True, WHITE)])
text(s, 1.0, 3.25, 11.3, 0.6, [("Trách nhiệm của mỗi chúng ta — vì một hành tinh xanh", 24, False, RGBColor(0xD7F2DB))])
text(s, 1.0, 6.35, 11.3, 0.5, [("Rem Agent  •  v3.26.1  •  2026", 14, False, RGBColor(0x9AD0A4))])

# ══════════ 2. MỤC LỤC ══════════
s = slide()
bg(s)
header(s, "Nội dung", "Chúng ta sẽ cùng tìm hiểu")
items = [
    ("01", "Môi trường là gì?", "vì sao nó quan trọng với sự sống"),
    ("02", "Các vấn đề cấp bách", "không khí, nước, rác thải nhựa"),
    ("03", "Biến đổi khí hậu", "thách thức toàn cầu"),
    ("04", "Hành động mỗi ngày", "những việc nhỏ tạo khác biệt lớn"),
    ("05", "Lời kêu gọi", "chung tay hành động ngay hôm nay"),
]
for i, (num, lab, sub) in enumerate(items):
    col = i % 2
    row = i // 2
    chip(s, 0.7 + col * 6.3, 1.7 + row * 1.95, 5.9, 1.7, num, lab, sub)

# ══════════ 3. MÔI TRƯỜNG LÀ GÌ ══════════
s = slide()
bg(s)
header(s, "Khái niệm", "Môi trường là gì")
bullets(s, 0.8, 1.6, 11.7, 3.2, [
    ("Môi trường là tất cả những gì bao quanh chúng ta", "không khí, nước, đất, rừng, động thực vật và con người."),
    ("Cung cấp tài nguyên sống còn", "không khí để thở, nước để uống, đất để trồng trọt, hệ sinh thái cân bằng."),
    ("Nơi tiếp nhận chất thải", "môi trường xử lý rác tự nhiên — nhưng bị quá tải khi xả thải vô tội vạ."),
])
boxes = [("Không khí", "trong lành để sống khỏe"), ("Nước", "nguồn sự sống"),
         ("Đất & rừng", "nền tảng lương thực"), ("Đa dạng sinh học", "cân bằng hệ sinh thái")]
for i, (t, d) in enumerate(boxes):
    col = i % 4
    shp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.7 + col * 3.1, 5.15, 2.9, 1.7, GREEN_M)
    text(s, 0.7 + col * 3.1 + 0.2, 5.35, 2.5, 1.3, [
        (t, 16, True, WHITE),
        (d, 11, False, RGBColor(0xE0F2E2))], PP_ALIGN.CENTER)

# ══════════ 4. Ô NHIỄM KHÔNG KHÍ ══════════
s = slide()
bg(s)
header(s, "Vấn đề 1", "Ô nhiễm không khí")
color = [RGBColor(0xC8B400), RGBColor(0xD84315), RGBColor(0xE64A19), RGBColor(0xB71C1C)]
stats = [("Khói bụi", "giao thông & nhà máy"), ("CO₂", "chất thải công nghiệp"),
         ("PM2.5", "hạt mịn xâm nhập phổi"), ("Hóa chất", "vô cơ & hữu cơ độc hại")]
for i, (t, d) in enumerate(stats):
    x = 0.7 + (i % 2) * 6.2
    y = 1.7 + (i // 2) * 2.1
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 5.9, 1.85, color[i], alpha=0.12, line=color[i])
    text(s, x + 0.35, y + 0.3, 5.2, 1.3, [
        (t, 20, True, color[i]),
        (d, 13, False, TEXT),
        ("“Hít thở không khí sạch là quyền cơ bản — hãy bảo vệ nó.”", 11, False, GRAY)])
bullets(s, 0.8, 6.0, 11.7, 1.2, [
    ("Mỗi năm, ô nhiễm không khí ảnh hưởng đến hàng triệu người, đặc biệt tại các đô thị lớn.", "")])

# ══════════ 5. Ô NHIỄM NƯỚC ══════════
s = slide()
bg(s)
header(s, "Vấn đề 2", "Ô nhiễm nguồn nước")
waves = [("Chất thải công nghiệp", "đổ thẳng ra sông hồ"), ("Thuốc trừ sâu", "từ nông nghiệp"),
         ("Rác thải sinh hoạt", "không được xử lý"), ("Dầu tràn", "tai nạn giao thông thủy")]
for i, (t, d) in enumerate(waves):
    col = i % 2; row = i // 2
    x = 0.7 + col * 6.2; y = 1.7 + row * 1.95
    c = RGBColor(0x0277BD)
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 5.9, 1.7, c, alpha=0.10, line=c)
    text(s, x + 0.35, y + 0.32, 5.2, 1.1, [
        (t, 18, True, c),
        (d, 13, False, TEXT)])
shape(s, MSO_SHAPE.RECTANGLE, 0.8, 5.7, 0.09, 1.8, RGBColor(0x0277BD))
text(s, 1.15, 5.7, 11.4, 1.5, [
    ("“Nước là sự sống — mỗi giọt nước bị ô nhiễm là một phần sự sống mất đi.”", 18, True, GREEN_D),
    ("Hạn chế xả thải, xử lý nước thải đúng quy trình, bảo vệ đầu nguồn.", 14, False, GRAY)])

# ══════════ 6. RÁC THẢI NHỰA ══════════
s = slide()
bg(s)
header(s, "Vấn đề 3", "Rác thải nhựa")
shape(s, MSO_SHAPE.OVAL, 0.8, 1.8, 2.4, 2.4, GREEN_D)
text(s, 0.8, 2.55, 2.4, 1.0, [("~430", 40, True, WHITE), ("triệu tấn nhựa/năm", 11, False, WHITE)], PP_ALIGN.CENTER)
bullets(s, 3.6, 1.8, 9.0, 3.4, [
    ("Nhựa phân hủy rất lâu", "túi nilon: 20–1000 năm mới phân hủy trong tự nhiên."),
    ("Vào đại dương", "hàng triệu tấn chảy ra biển mỗi năm, rác thải nhựa đe dọa sinh vật biển."),
    ("Vào cơ thể chúng ta", "vi nhựa xuất hiện trong nước, thực phẩm và cả không khí."),
    ("Giải pháp đơn giản", "giảm dùng nhựa 1 lần: túi vải, bình nước, ống hút inox.")])
text(s, 0.8, 6.25, 11.7, 0.6, [("3R thần thánh:  Reduce • Reuse • Recycle", 20, True, GREEN_L)])

# ══════════ 7. BIẾN ĐỔI KHÍ HẬU ══════════
s = slide()
bg(s)
header(s, "Thách thức", "Biến đổi khí hậu")
rows = [("Nhiệt độ tăng", "hiện tượng El Niño, phát thải CO₂ từ đốt nhiên liệu hóa thạch"),
        ("Băng tan & nước biển dâng", "đe dọa các vùng ven biển, đảo quốc"),
        ("Thời tiết cực đoan", "bão lũ, hạn hán, nắng nóng kỷ lục diễn ra thường xuyên hơn")]
for i, (t, d) in enumerate(rows):
    y = 1.75 + i * 1.75
    shape(s, MSO_SHAPE.RECTANGLE, 0.85, y, 0.09, 1.25, GOLD)
    text(s, 1.2, y, 11.3, 1.4, [
        (t, 20, True, GREEN_D),
        (d, 14, False, TEXT)])
text(s, 0.8, 6.55, 11.7, 0.5, [("COP và các hiệp định toàn cầu kêu gọi giảm phát thải ròng bằng 0.", 13, False, GRAY)])

# ══════════ 8. HÀNH ĐỘNG MỖI NGÀY ══════════
s = slide()
bg(s)
header(s, "Giải pháp", "Những hành động nhỏ, ảnh hưởng lớn")
acts = [("1", "Tắt điện khi không dùng"), ("2", "Phân loại rác tại nguồn"),
        ("3", "Dùng bình nước, túi vải"), ("4", "Trồng cây xanh quanh nhà"),
        ("5", "Đi bộ / xe buýt / xe đạp"), ("6", "Nhắc nhau bảo vệ môi trường")]
for i, (n, t) in enumerate(acts):
    col = i % 3; row = i // 3
    x = 0.7 + col * 4.15; y = 1.8 + row * 2.3
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 3.9, 2.0, GREEN_M)
    circ = shape(s, MSO_SHAPE.OVAL, x + 0.3, y + 0.3, 1.4, 1.4, GOLD)
    text(s, x + 0.62, y + 0.5, 0.8, 0.8, [(n, 24, True, GREEN_D)], PP_ALIGN.CENTER)
    text(s, x + 1.95, y + 0.45, 1.8, 1.2, [(t, 15, True, WHITE)])

# ══════════ 9. LỜI KÊU GỌI ══════════
s = slide()
bg(s, GREEN_D)
shape(s, MSO_SHAPE.OVAL, 9.8, 4.9, 3.2, 3.2, GREEN_M, alpha=0.4)
shape(s, MSO_SHAPE.OVAL, 10.6, 5.8, 1.9, 1.9, GOLD, alpha=0.3)
text(s, 1.2, 2.0, 11.0, 1.0, [("MỖI HÀNH ĐỘNG NHỎ ĐỀU CÓ GIÁ TRỊ", 40, True, WHITE)], PP_ALIGN.CENTER)
text(s, 1.2, 3.2, 11.0, 1.6, [
    ("“Trái đất không phải là tài sản ta thừa hưởng từ cha ông,", 22, False, RGBColor(0xD7F2DB)),
    (" mà là món nợ ta vay từ con cháu.”", 22, False, RGBColor(0xD7F2DB))], PP_ALIGN.CENTER)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 4.65, 5.3, 4.0, 1.0, GOLD)
text(s, 4.65, 5.55, 4.0, 0.6, [("BẮT ĐẦU NGAY HÔM NAY", 20, True, GREEN_D)], PP_ALIGN.CENTER)

# ══════════ 10. CẢM ƠN ══════════
s = slide()
bg(s)
shape(s, MSO_SHAPE.OVAL, 10.9, -1.0, 3.6, 3.6, GREEN_L, alpha=0.5)
text(s, 0.8, 2.3, 11.7, 1.0, [("CẢM ƠN! 💚", 54, True, GREEN_D)], PP_ALIGN.CENTER)
text(s, 0.8, 3.6, 11.7, 0.7, [("Trái đất xanh — chung tay vì thế hệ tương lai", 22, False, GREEN_M)], PP_ALIGN.CENTER)
text(s, 0.8, 6.2, 11.7, 0.5, [("Rem Agent • Bộ giáo trình phổ biến về môi trường", 13, False, GRAY)], PP_ALIGN.CENTER)

prs.save(OUT)
print("SAVED", OUT, os.path.getsize(OUT), "bytes")